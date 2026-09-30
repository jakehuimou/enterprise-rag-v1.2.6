"""
结果导出层 - 问答内容的 Word / Excel 导出
==========================================
把对话回答（Markdown 文本）渲染成可下载的办公文档：

  - ``build_docx``  ：Markdown -> .docx（标题、段落、列表、引用、代码块、表格）
  - ``build_xlsx``  ：抽取 Markdown 表格 -> .xlsx（多表分 Sheet，表头冻结 + 样式）
  - ``has_table``   ：快速判断一段文本里是否含 Markdown 表格（供前端决定按钮是否可用）

设计要点：
  - 不依赖 ``markdown`` / ``pandoc`` 等外部转换器，纯 python-docx + openpyxl 实现，
    保证离线可用；
  - Word 导出保留「参考来源」等全部文字；Excel 只取表格内容（表格才是数据）。

输出均为 ``bytes``，由调用方决定落盘或直接作为 HTTP 响应体返回。
"""
import io
import re
from datetime import datetime
from typing import List, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Markdown 解析辅助
# ---------------------------------------------------------------------------
# 表格分隔行，例如 | --- | :---: | ---: |
_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)+\|?\s*$")
# 代码块围栏
_FENCE_RE = re.compile(r"^\s*```")
# 「参考来源」段落（前端在答案末尾拼接的引用块）
_SOURCE_BLOCK_RE = re.compile(r"\n{2,}【参考来源】[\s\S]*$")
# 有序列表
_OL_RE = re.compile(r"^\s*(\d+)[.、)]\s+(.*)$")
# 无序列表
_UL_RE = re.compile(r"^\s*[-*+]\s+(.*)$")
# 行内代码 / 粗体
_INLINE_TOKEN_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def strip_source_block(text: str) -> str:
    """去掉答案末尾的「【参考来源】…」引用块（导出表格时不需要）。"""
    return _SOURCE_BLOCK_RE.sub("", text or "").strip()


def has_source_block(text: str) -> bool:
    """判断文本是否带有「【参考来源】」引用块。"""
    return bool(_SOURCE_BLOCK_RE.search(text or ""))


def _split_row(line: str) -> List[str]:
    """把一行 Markdown 表格拆成单元格列表（支持 \\| 转义）。"""
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    # 先保护转义竖线，再按竖线切分
    cells = re.split(r"(?<!\\)\|", s)
    return [c.replace("\\|", "|").strip() for c in cells]


def _clean_cell(cell: str) -> str:
    """去掉单元格里的 Markdown 标记（粗体 / 行内代码 / 链接）。"""
    c = re.sub(r"\*\*([^*]+)\*\*", r"\1", cell)
    c = re.sub(r"`([^`]+)`", r"\1", c)
    c = re.sub(r"\[([^\]]*)\]\(([^)]*)\)", r"\1", c)
    return c.strip()


def extract_tables(text: str) -> List[dict]:
    """
    抽取文本中的所有 Markdown 表格。

    返回列表，每项为 ``{"caption": str, "rows": [[cell, ...], ...]}``：
      - ``rows[0]`` 为表头；
      - ``caption`` 为表格正上方最近的一行非空文本（用于 Excel Sheet 名），可能为空。
    """
    lines = (text or "").splitlines()
    tables: List[dict] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        is_header = ("|" in line.strip()) and i + 1 < len(lines) and bool(_TABLE_SEP_RE.match(lines[i + 1]))
        if not is_header:
            i += 1
            continue
        header = _split_row(line)
        body: List[List[str]] = []
        j = i + 2
        while j < len(lines) and "|" in lines[j] and lines[j].strip():
            body.append(_split_row(lines[j]))
            j += 1
        width = len(header)
        rows = [header] + [
            (r + [""] * width)[:width] for r in body
        ]
        rows = [[_clean_cell(c) for c in r] for r in rows]
        # 表标题：表格上方最近的一行非空文本（去掉 #、>、* 等 Markdown 标记），
        # 用作 Excel Sheet 名，例如「一、方案对比」。
        caption = ""
        k = i - 1
        while k >= 0:
            cand = lines[k].strip()
            if cand:
                if "|" not in cand:
                    caption = re.sub(r"^[#>\*\s]+|[\*\s]+$", "", cand).strip()
                break
            k -= 1
        tables.append({"caption": caption, "rows": rows})
        i = j
    return tables


def has_table(text: str) -> bool:
    """判断文本中是否含至少一个 Markdown 表格（含表头分隔行）。"""
    return bool(extract_tables(text))


# ---------------------------------------------------------------------------
# Word 导出
# ---------------------------------------------------------------------------
_BODY_FONT = "微软雅黑"
_MONO_FONT = "Consolas"


def _apply_font(run, name: str = _BODY_FONT, size: int = None,
                bold: bool = False, color: Optional[tuple] = None):
    """给 run 设置中英文字体（中文必须显式写 w:eastAsia，否则 Word 会回落默认宋体）。"""
    run.font.name = name
    run.font.bold = bold
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def _add_runs(paragraph, text: str, base_size: int = 11, bold: bool = False):
    """写入带行内样式的文本（**粗体** / `代码`）。"""
    for token in _INLINE_TOKEN_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            _apply_font(paragraph.add_run(token[2:-2]), _BODY_FONT, base_size, True)
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            _apply_font(paragraph.add_run(token[1:-1]), _MONO_FONT, base_size - 1)
        else:
            _apply_font(paragraph.add_run(token), _BODY_FONT, base_size, bold)


def _add_table(doc: Document, rows: List[List[str]]):
    """把二维数据写入 Word 表格。"""
    if not rows:
        return
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci in range(width):
            cell = cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            _add_runs(p, row[ci] if ci < len(row) else "", base_size=10, bold=(ri == 0))


def build_docx(title: str, content: str) -> bytes:
    """
    把 Markdown 文本渲染为 .docx 字节流。

    :param title: 文档标题（一级标题）
    :param content: 回答正文（Markdown）
    """
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = _BODY_FONT
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), _BODY_FONT)

    if title:
        h = doc.add_heading("", level=0)
        h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _apply_font(h.add_run(title), _BODY_FONT, 20, True, (0x1F, 0x2A, 0x24))

    in_code = False
    buf: List[str] = []
    lines = (content or "").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        raw = line.rstrip()

        # 代码块
        if _FENCE_RE.match(raw):
            if in_code:
                p = doc.add_paragraph()
                _apply_font(p.add_run("\n".join(buf)), _MONO_FONT, 9, color=(0x33, 0x33, 0x33))
                buf, in_code = [], False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            buf.append(raw)
            i += 1
            continue

        # 表格
        if i + 1 < len(lines) and "|" in raw and _TABLE_SEP_RE.match(lines[i + 1]):
            rows = [_split_row(raw)]
            j = i + 2
            while j < len(lines) and "|" in lines[j] and lines[j].strip():
                rows.append(_split_row(lines[j]))
                j += 1
            rows = [[_clean_cell(c) for c in r] for r in rows]
            _add_table(doc, rows)
            doc.add_paragraph()
            i = j
            continue

        stripped = raw.strip()
        if not stripped:
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            text = re.sub(r"\*\*", "", m.group(2)).strip()
            h = doc.add_heading("", level=min(len(m.group(1)), 4))
            _apply_font(h.add_run(text), _BODY_FONT, 16 - 2 * len(m.group(1)), True)
            i += 1
            continue

        # 分隔线
        if re.match(r"^\s*([-*_])\s*\1\s*\1[\s\-*_]*$", stripped):
            i += 1
            continue

        # 引用
        if stripped.startswith(">"):
            text = stripped.lstrip("> ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            _apply_font(p.add_run("　"), _BODY_FONT, 11)
            _add_runs(p, text, base_size=10)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x55, 0x66, 0x5C)
            i += 1
            continue

        # 有序 / 无序列表
        m = _OL_RE.match(raw)
        if m:
            p = doc.add_paragraph(style="List Number")
            _add_runs(p, m.group(2))
            i += 1
            continue
        m = _UL_RE.match(raw)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, m.group(1))
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        _add_runs(p, stripped)
        i += 1

    if in_code and buf:  # 未闭合的代码块也要落盘
        p = doc.add_paragraph()
        _apply_font(p.add_run("\n".join(buf)), _MONO_FONT, 9)

    foot = doc.add_paragraph()
    _apply_font(
        foot.add_run(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}　·　企业文档知识库 RAG"),
        _BODY_FONT, 8, color=(0x99, 0x99, 0x99),
    )

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


# ---------------------------------------------------------------------------
# Excel 导出
# ---------------------------------------------------------------------------
_HEADER_FILL = PatternFill("solid", fgColor="2F5D50")
_HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
_BODY_FONT_X = Font(name="微软雅黑", size=10.5)
_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _safe_sheet_name(name: str, used: set, fallback: str) -> str:
    """清洗 Excel Sheet 名（去非法字符、限 31 字符、去重）。"""
    clean = re.sub(r"[\[\]:*?/\\]", " ", (name or "").strip())
    clean = re.sub(r"\s+", " ", clean)[:28].strip()
    if not clean:
        clean = fallback
    base, n = clean, 1
    while clean in used:
        n += 1
        clean = f"{base}_{n}"[:31]
    used.add(clean)
    return clean


def _to_number(v: str):
    """把「12,345.6」/「12%」这类文本转成数值，失败则原样返回。"""
    s = (v or "").strip().replace(",", "").replace("，", "")
    if not s:
        return v
    pct = s.endswith("%")
    if pct:
        s = s[:-1]
    try:
        num = float(s)
    except Exception:
        return v
    if pct:
        return num / 100.0
    return int(num) if num.is_integer() else num


def build_xlsx(content: str, sheet_prefix: str = "表格") -> bytes:
    """
    抽取 Markdown 表格并生成 .xlsx 字节流；一个表格一个 Sheet。

    :raises ValueError: 文本中没有任何 Markdown 表格
    """
    tables = extract_tables(strip_source_block(content))
    if not tables:
        raise ValueError("当前回答中没有检测到表格数据，无法导出 Excel。")

    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    for idx, tb in enumerate(tables, start=1):
        rows = tb["rows"]
        title = _safe_sheet_name(tb.get("caption", ""), used, f"{sheet_prefix}{idx}")
        ws = wb.create_sheet(title=title)

        width = max(len(r) for r in rows)
        for ri, row in enumerate(rows):
            for ci in range(width):
                raw = row[ci] if ci < len(row) else ""
                cell = ws.cell(row=ri + 1, column=ci + 1,
                               value=(raw if ri == 0 else _to_number(raw)))
                # 「8.5%」这类百分数转成 0.085 后补上百分比格式，避免显示成 0.085
                if ri > 0 and isinstance(cell.value, (int, float)) and raw.strip().endswith("%"):
                    cell.number_format = "0.00%"
                cell.border = _BORDER
                cell.alignment = Alignment(
                    horizontal="center" if ri == 0 else "left",
                    vertical="center", wrap_text=True,
                )
                cell.font = _HEADER_FONT if ri == 0 else _BODY_FONT_X
                if ri == 0:
                    cell.fill = _HEADER_FILL
        ws.freeze_panes = "A2"
        if len(rows) > 1:
            ws.auto_filter.ref = f"A1:{get_column_letter(width)}{len(rows)}"
        # 列宽自适应（按最长内容估算，限 8~42）
        for ci in range(width):
            longest = max(
                (len(str(rows[ri][ci])) if ci < len(rows[ri]) else 0)
                for ri in range(len(rows))
            )
            ws.column_dimensions[get_column_letter(ci + 1)].width = max(
                8, min(42, longest * 1.6 + 4)
            )
        ws.row_dimensions[1].height = 24

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def export_filename(kind: str = "word") -> str:
    """生成带时间戳的导出文件名。"""
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    ext = "docx" if kind == "word" else "xlsx"
    return f"问答导出_{stamp}.{ext}"
